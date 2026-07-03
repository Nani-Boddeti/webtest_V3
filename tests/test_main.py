"""Tests for the /generate and /health endpoints."""

import pytest
from httpx import AsyncClient
from io import BytesIO
from docx import Document


@pytest.mark.anyio
async def test_health(client: AsyncClient):
    resp = await client.get("/health")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


@pytest.mark.anyio
async def test_generate_returns_docx(client: AsyncClient):
    resp = await client.post(
        "/generate", json={"topic": "AI website maintenance agent"}
    )
    assert resp.status_code == 200
    assert (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        in resp.headers["content-type"]
    )
    assert "attachment" in resp.headers["content-disposition"]


@pytest.mark.anyio
async def test_generate_docx_content_valid(client: AsyncClient):
    """Verify the .docx contains the topic and all five required sections."""
    resp = await client.post("/generate", json={"topic": "Test topic"})
    assert resp.status_code == 200

    doc = Document(BytesIO(resp.content))
    paragraphs = [p.text for p in doc.paragraphs]
    full_text = "\n".join(paragraphs)

    # Topic must appear in the report
    assert "Test topic" in full_text

    # All five required sections must be present
    assert "5-Point Summary" in full_text
    assert "Challenges to Resolve" in full_text
    assert "Optimal Model Recommendation" in full_text
    assert "Model Guardrails" in full_text
    assert "Implementation Guidelines" in full_text


@pytest.mark.anyio
async def test_generate_rejects_empty_topic(client: AsyncClient):
    resp = await client.post("/generate", json={"topic": ""})
    assert resp.status_code == 422


@pytest.mark.anyio
async def test_generate_rejects_missing_topic(client: AsyncClient):
    resp = await client.post("/generate", json={})
    assert resp.status_code == 422


@pytest.mark.anyio
async def test_generate_docx_has_expected_structure(client: AsyncClient):
    """Minimal structural checks: verify the document has headings and body text."""
    resp = await client.post("/generate", json={"topic": "Structure check"})
    assert resp.status_code == 200

    doc = Document(BytesIO(resp.content))

    # Should have at least one heading
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) >= 1

    # Should have non-empty body paragraphs beyond just the title
    non_empty = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert len(non_empty) > 3
