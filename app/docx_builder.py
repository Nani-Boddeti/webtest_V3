"""
Convert structured document text into a .docx byte stream using python-docx.

This module is intentionally small — it wraps python-docx calls so the
rest of the application doesn't need to know about the document library.
"""

from io import BytesIO

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_docx(topic: str, document_text: str) -> bytes:
    """Return a .docx file as raw bytes for the given topic and text.

    *document_text* is expected to be plain text with section headings
    delimited by `##`.  The builder creates proper Word headings and
    body paragraphs so the result looks professional.
    """
    doc = Document()

    # --- Default style tweaks ---
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    # --- Title ---
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(f"Research Report: {topic}")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(12)

    # --- Parse lines ---
    for line in document_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue

        # Skip the ASCII-art separator
        if stripped.startswith("===") or stripped.startswith("---"):
            continue

        if stripped.startswith("## "):
            heading = doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            heading = doc.add_heading(stripped[2:], level=1)
        else:
            doc.add_paragraph(stripped)

    # --- Serialise ---
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
